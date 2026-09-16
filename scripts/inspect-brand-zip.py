import zipfile,pathlib,json,logging
from pypdf import PdfReader
import openpyxl
logging.getLogger('pdfminer').setLevel(logging.ERROR)
out=pathlib.Path('tmp/brands-20260916');out.mkdir(parents=True,exist_ok=True)
with zipfile.ZipFile(r'C:\Users\ADMIN\Downloads\BÁO GIÁ.zip') as z:
 for f in z.infolist():
  if f.is_dir():continue
  dest=out/pathlib.PurePosixPath(f.filename).name
  dest.write_bytes(z.read(f))
summary=[]
for p in out.iterdir():
 if p.suffix.lower()=='.pdf':
  r=PdfReader(p); pages=[page.extract_text() or '' for page in r.pages]
  (out/(p.stem+'.txt')).write_text('\n\n'.join('=== PAGE '+str(i+1)+' ===\n'+s for i,s in enumerate(pages)),encoding='utf-8')
  summary.append({'file':p.name,'pages':[len(s) for s in pages]})
 if p.suffix.lower()=='.xlsx':
  w=openpyxl.load_workbook(p,data_only=True)
  rows={s.title:[{'row':i,'cells':{openpyxl.utils.get_column_letter(j):str(v) if not isinstance(v,(str,int,float,bool,type(None))) else v for j,v in enumerate(row,1) if v is not None}} for i,row in enumerate(s.values,1) if any(v is not None for v in row)] for s in w}
  (out/(p.stem+'.json')).write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
  summary.append({'file':p.name,'sheets':[{ 'name':s.title,'rows':s.max_row,'cols':s.max_column,'images':len(s._images)} for s in w]})
print(json.dumps(summary,ensure_ascii=False,indent=2))
