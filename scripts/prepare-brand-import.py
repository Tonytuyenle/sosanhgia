import pathlib,json,re,hashlib,collections,logging,unicodedata,io
import openpyxl,pdfplumber
from pypdf import PdfReader
from PIL import Image
logging.getLogger('pdfminer').setLevel(logging.ERROR)
B=pathlib.Path('tmp/brands-20260916');A=B/'assets';A.mkdir(exist_ok=True)
rows=[];sources=[]
def clean(v):return re.sub(r'\s+',' ',str(v or '')).strip().replace('ƣ','ư').replace('Ƣ','Ư')
def number(v):
 if v is None or v=='':return None
 if isinstance(v,(int,float)):return round(v)
 s=re.sub(r'[\s₫()]','',str(v))
 if re.fullmatch(r'\d{1,3}(?:[.,]\d{3})+',s):return int(re.sub(r'[.,]','',s))
 if re.fullmatch(r'\d+\.0+',s):return int(float(s))
 if re.fullmatch(r'\d+',s):return int(s)
 return None
def source(p,slug):
 dest=slug+p.suffix.lower();(A/dest).write_bytes(p.read_bytes());sources.append({'file':p.name,'asset':dest,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()});return '/api/assets/brands-20260916/'+dest
def category(name):
 n=name.lower()
 if 'combo' in n:return 'Combo sản phẩm'
 if any(x in n for x in ['vung','xửng','đũa','thớt','dao','muôi','dụng cụ','hộp','túi','nạo']):return 'Dụng cụ nhà bếp'
 if any(x in n for x in ['nước giặt','nước rửa']):return 'Chăm sóc nhà cửa'
 if any(x in n for x in ['điện','máy','bếp từ','quạt','bếp điện','cây nước','siêu tốc','nồi nấu chậm','nồi lẩu','nồi cơm']):return 'Đồ điện'
 if 'chảo' in n:return 'Chảo'
 if 'áp suất' in n:return 'Nồi tăng áp'
 if any(x in n for x in ['nồi','quánh']):return 'Nồi & bộ nồi'
 if any(x in n for x in ['bình','cốc','ấm']):return 'Ấm & bình'
 return 'Gia dụng khác'
def add(brand,code,name,loc,p,prices,desc='',extra='',image=None,warnings=None):
 item={'brand':brand,'code':clean(code).upper(),'model':clean(code),'name':clean(name),'category':category(clean(name)),'description':str(desc or '').strip(),'source':p.name+' — '+loc+' — người dùng cung cấp','priceSource':sourceurls[p.name],'extra':extra,'supplier':brand,**{k:v for k,v in prices.items() if v is not None}}
 if image:item['image']=image
 row={'input':item,'source':{'file':p.name,'location':loc},'warnings':warnings or []};rows.append(row);return row
def xvalue(s,r,c):
 v=s.cell(r,c).value
 if v is not None:return v
 for m in s.merged_cells.ranges:
  if m.min_row<=r<=m.max_row and m.min_col<=c<=m.max_col:return s.cell(m.min_row,m.min_col).value
 return None
def ximages(s,slug):
 out=[]
 for i,img in enumerate(s._images):
  if not hasattr(img.anchor,'_from'):continue
  r=img.anchor._from.row+1;end=img.anchor.to.row+1 if hasattr(img.anchor,'to') else r
  im=Image.open(io.BytesIO(img._data())).convert('RGB');im.thumbnail((1000,1000));dest=f'{slug}-img-{i+1}.jpg';im.save(A/dest,quality=90)
  out.append((r,end,'/api/assets/brands-20260916/'+dest))
 return out
def xi(images,row):
 found=[i for i in images if i[0]<=row<=i[1]]
 return min(found,key=lambda i:abs(i[0]-row))[2] if found else None
def pdfimage(page,reader,box,slug):
 if not box:return None
 found=[]
 for im in page.images:
  area=max(0,min(box[2],im['x1'])-max(box[0],im['x0']))*max(0,min(box[3],im['bottom'])-max(box[1],im['top']))
  if area>4:found.append((area,im))
 if not found:return None
 im=max(found,key=lambda x:x[0])[1]
 try:
  obj=reader.pages[page.page_number-1].images['/'+im['name']];pic=Image.open(io.BytesIO(obj.data)).convert('RGB');pic.thumbnail((1000,1000));dest=slug+'-'+hashlib.sha256(obj.data).hexdigest()[:12]+'.jpg';pic.save(A/dest,quality=90);return '/api/assets/brands-20260916/'+dest
 except Exception:return None
def cellbox(t,i,c):
 box=t['cells'][i][c]
 if box:return box
 ref=t['cells'][i][6] or t['cells'][i][0];y=(ref[1]+ref[3])/2
 return next((r[c] for r in t['cells'] if r[c] and r[c][1]<=y<=r[c][3]),None)
sourceurls={}
for p in B.iterdir():
 if p.suffix in ['.pdf','.xlsx']:
  slug=next(x for x in ['hare','engler','fynal','haatz','jiwon','morico','kailer'] if x in p.name.lower())
  sourceurls[p.name]=source(p,slug)

for p in B.glob('*.xlsx'):
 s=openpyxl.load_workbook(p,data_only=True).worksheets[0]
 slug=next(x for x in ['hare','fynal','kailer'] if x in p.name.lower());images=ximages(s,slug)
 if slug=='hare':
  brand='Hare'
  for r in range(13,62):
   if s.cell(r,1).value=='YOAKE':brand='Yoake'
   if not s.cell(r,3).value:continue
   vals={k:number(s.cell(r,c).value) for k,c in [('distributorPrice',6),('minimum',7),('listPrice',8)]}
   w=[]
   if vals['minimum']>vals['listPrice']:w.append('Giá bán tối thiểu cao hơn giá niêm yết trong nguồn; cần xác minh.');vals.pop('minimum')
   row=add(brand,s.cell(r,3).value,s.cell(r,2).value,f'Sheet {s.title}, dòng {r}',p,vals,s.cell(r,5).value,'Áp dụng từ 01/02/2026. Chiết khấu doanh số năm: 1% từ 1,5 tỷ; 2% từ 2 tỷ, chưa tự trừ vào giá.\nGiá bán tối thiểu gốc: '+str(s.cell(r,7).value),xi(images,r),w);row['input']['quoteDate']='2026-02-01'
 if slug=='fynal':
  for r in range(8,30):
   rawcode=clean(s.cell(r,5).value);code=re.sub(r'\s*\(.*','',rawcode).strip();name=xvalue(s,r,2)
   if not code:continue
   if r in [25,26]:name=clean(name)+' '+code
   w=[]
   if r in [12,13]:w.append('Mã trong tên sản phẩm khác cột Mã sản phẩm; giữ nguyên cả hai để xác minh.')
   margin=s.cell(r,8).value
   if isinstance(margin,str) and margin.startswith('#'):w.append('Nguồn có lỗi công thức biên lợi nhuận '+margin+'; không nhập thành số.')
   add('Fynal',code,name,f'Sheet {s.title}, dòng {r}',p,{'distributorPrice':number(s.cell(r,6).value),'minimum':number(s.cell(r,7).value)},xvalue(s,r,4),'Mã/quy cách gốc: '+rawcode+'\nTồn kho theo báo giá: '+clean(s.cell(r,9).value)+'\nBiên độ lợi nhuận gốc: '+str(margin)+'\nGiá tại kho; chi phí đóng gói đơn lẻ tính riêng. Ngày ký chưa điền (mẫu năm 2025).',xi(images,r),w)['input']['stock']=clean(s.cell(r,9).value)
 if slug=='kailer':
  for r in range(3,68):
   brand=clean(s.cell(r,3).value);name=clean(s.cell(r,5).value)
   if not brand or not name:continue
   brand={'KAILER':'Kailer','BUCOOK':'Bucook','BUHAIBER':'Buhaiber','KAISA VILLA':'Kaisa Villa'}[brand]
   codefield=str(s.cell(r,4).value or '');codes=[clean(x) for x in codefield.splitlines() if clean(x) and not clean(x).startswith('(')] or [f'KAILER-NOCODE-R{r}']
   wp=s.cell(r,7).value;lp=s.cell(r,6).value
   lines=lambda v:[x.strip() for x in str(v or '').splitlines() if x.strip()]
   wl=lines(wp);ll=lines(lp)
   for j,code in enumerate(codes):
    warnings=[];prices={};size=''
    raww=wl[j] if len(wl)==len(codes) else wp
    m=re.fullmatch(r'(?:Size\s*)?(\d+)\s*(cm|L)?\s*[:-]\s*([\d.,]+)',str(raww).strip(),re.I)
    if m:size=m[1]+(' l' if m[2] and m[2].upper()=='L' else ' cm');v=number(m[3])
    elif r==56:v=55000
    else:v=number(raww)
    if v is not None:
     if v<10000:warnings.append('Giá buôn gốc '+str(raww)+' chưa rõ đơn vị; chưa đưa vào cột giá VND.')
     else:prices['wholesalePrice']=v
    rawl=ll[j] if len(ll)==len(codes) else lp
    if len(codes)>1 and len(ll)!=len(codes):
     if lp is not None:warnings.append('Số mức giá niêm yết không khớp số biến thể; giữ nguyên nguồn để xác minh.')
    else:
     v=number(rawl)
     if v is not None:
      if v<10000:warnings.append('Giá niêm yết gốc '+str(rawl)+' chưa rõ đơn vị; chưa đưa vào cột giá VND.')
      else:prices['listPrice']=v
    if not codefield:warnings.append('Nguồn không có mã; dùng mã nội bộ theo dòng bảng giá.')
    if not prices:warnings.append('Chưa có mức giá VND xác định được trong nguồn.')
    item=add(brand,code,name+(' — '+size if len(codes)>1 else ''),f'Sheet {s.title}, dòng {r}, biến thể {j+1}',p,prices,name,'Giá niêm yết gốc: '+str(lp)+'\nGiá buôn gốc: '+str(wp)+'\nMã/quy cách gốc: '+codefield+'\nGhi chú: '+clean(s.cell(r,9).value),xi(images,r),warnings)['input']
    item['stock']=clean(s.cell(r,8).value)
    if size:item['dimensions' if size.endswith('cm') else 'capacity']=size
    if r==56:item['moq']=150
    if r==18:item['minimum']=1190000

for slug in ['engler','jiwon','haatz']:
 p=next(p for p in B.glob('*.pdf') if slug in p.name.lower());reader=PdfReader(p)
 tables=json.loads((B/(p.stem+'-tables.json')).read_text(encoding='utf-8'))
 with pdfplumber.open(p) as doc:
  for t in tables:
   page=doc.pages[t['page']-1];lastcode=''
   for i,r in enumerate(t['rows']):
    if slug!='haatz' and not str(r[0]).isdigit():continue
    if slug=='haatz' and number(r[6]) is None:continue
    ref=t['cells'][i][6 if slug=='haatz' else 5];y0,y1=ref[1],ref[3]
    def cell(c,crop=False):
     box=cellbox(t,i,c)
     if not box:return ''
     if crop:box=(box[0],max(box[1],y0),box[2],min(box[3],y1))
     return page.crop(box).extract_text(x_tolerance=1) or ''
    picturebox=cellbox(t,i,1);pic=pdfimage(page,reader,picturebox,slug)
    loc=f'Trang {t["page"]}, dòng bảng {i+1}, STT {r[0] or "biến thể"}'
    if slug=='engler':
     code=clean(cell(2));name=clean(r[3]);w=[]
     if not code:code=f'ENGLER-P{t["page"]}-R{i+1}';w=['Nguồn không có mã; dùng mã nội bộ theo vị trí dòng.']
     add('Engler',code,name,loc,p,{'distributorPrice':number(r[5]),'listPrice':number(r[7]),'marketplacePrice':number(r[8]),'packingFee':number(r[9])},name,'Quy cách: '+clean(r[4])+'\nChiết khấu NPP theo nguồn: '+clean(r[6])+' (không trừ thêm vào giá NPP).',pic,w)
    if slug=='jiwon':
     name=clean(cell(2,True));code=re.sub(r'\s*-\s*','-',clean(r[3]));desc=cell(4,True);w=[]
     if int(r[0])==28:w.append('Bảo hành trong mô tả là 12 tháng, cột bảo hành ghi 2 năm; cần xác minh.')
     item=add('Jiwon',code,name,loc,p,{'bulk20Price':number(r[5]),'bulk50Price':number(r[6]),'bulk100Price':number(r[7]),'listPrice':number(r[8]),'websitePrice':number(r[9])},desc,'Giá theo số lượng: từ 20 / 50 / 100 cái. Mức từ 100 cái dành cho NPP/khách live lớn.\nBảo hành gốc: '+clean(r[10]),pic,w)['input']
     item['channel']='NPP / Khách live lớn; điều kiện theo số lượng'
    if slug=='haatz':
     code=clean(cell(2));name=clean(cell(3));spec=clean(r[5]);w=[]
     if t['page'] in [4,5,6,7]:
      name={4:'Chảo Yeon Inox 316 cạn lòng',5:'Chảo Yeon Inox 316 sâu lòng',6:'Chảo/nồi lẩu Yeon Inox 316',7:'Quánh Yeon Inox 316'}[t['page']]
      m=re.search(r'(\d+)\s*cm',spec)
      if m:name+=' '+m[1]+' cm'
      if t['page'] in [5,6,7] and code!='HYP312':name+=' — '+('kèm vung' if re.search(r'vung',spec,re.I) else 'không kèm vung')
     item=add('Haatz',code,name,loc,p,{'nppOnline':number(r[6]),'listPrice':number(r[7]),'retailPrice':number(r[8])},spec,'Áp dụng với nhà bán online từ tháng 1/2026.\nLợi nhuận theo báo giá: '+clean(r[9])+'\nTình trạng/ghi chú: '+clean(r[10])+'\nSố hàng/thùng: '+clean(cell(11)),pic,w)['input'];item['stock']=clean(r[10]);item['channel']='Nhà bán online'
     if code=='HTA628':item['extra']+='\nGiá mới chưa VAT theo ghi chú nguồn.'

p=next(p for p in B.glob('*.pdf') if 'Morico' in p.name)
for i,v in enumerate(json.loads(pathlib.Path('scripts/morico-reviewed.json').read_text(encoding='utf-8')),1):
 code,name,mall,retail=v[:4];w=[];prices={'mallPrice':mall,'retailPrice':retail}
 if i in [27,30]:
  w.append('Cột Đại lý/CTV/Shop thường ghi '+str(retail)+' đ, chỉ bằng 1/10 giá Mall; chờ xác minh, chưa dùng làm giá bán lẻ.');prices.pop('retailPrice')
 if i>=36:w.append('Combo có nhiều mã thành phần; mã COMBO là mã nội bộ để quản lý riêng.')
 item=add('Morico',code,name,f'Trang 1, STT {i}',p,prices,name,'Bảng giá sản phẩm ECOM 2026, áp dụng từ 01/09/2026 đến khi có thông báo mới.\nGiá Mall Morico: '+str(mall)+'\nGiá Đại lý, CTV, Shop thường: '+str(retail)+'\nGiá trên đã bao gồm chi phí vận chuyển.\n'+(v[4] if len(v)>4 else ''),warnings=w)['input'];item['quoteDate']='2026-09-01';item['channel']='ECOM: Mall Morico / Đại lý, CTV, Shop thường'

# Preserve repeated supplier codes while giving each distinct variant its own application code.
counts=collections.Counter((r['input']['brand'],r['input']['code']) for r in rows)
for row in rows:
 p=row['input'];key=(p['brand'],p['code'])
 if counts[key]>1:
  raw=p['code'];suffix=hashlib.sha256((p['brand']+'|'+row['source']['location']).encode()).hexdigest()[:6].upper();p['code']=raw+'-V'+suffix
  row['warnings'].append('Mã nhà cung cấp dùng chung nhiều dòng/biến thể; mã nội bộ thêm hậu tố, Model giữ mã gốc '+raw+'.')
 p['extra']+='\n'+'\n'.join(row['warnings'])
 assert p['name'] and p['code'],row
(B/'prepared.json').write_text(json.dumps({'sources':sources,'products':rows},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'counts':dict(collections.Counter(r['input']['brand'] for r in rows)),'warnings':sum(bool(r['warnings']) for r in rows),'images':sum(bool(r['input'].get('image')) for r in rows)},ensure_ascii=False))
